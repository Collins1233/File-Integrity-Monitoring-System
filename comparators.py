import difflib

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


# ==========================================================
# TEXT FILES
# ==========================================================

def read_text_file(file_path):
    """
    Reads a text file and returns its contents.
    """

    try:

        with open(file_path, "r", encoding="utf-8") as file:
            return file.readlines()

    except UnicodeDecodeError:

        try:

            with open(file_path, "r", encoding="latin-1") as file:
                return file.readlines()

        except Exception:

            return []

    except Exception:

        return []


def compare_text(old_file, new_file):
    """
    Compare two text files.
    """

    old_lines = read_text_file(old_file)

    new_lines = read_text_file(new_file)

    return list(

        difflib.unified_diff(

            old_lines,

            new_lines,

            fromfile="Baseline Version",

            tofile="Current Version",

            lineterm=""

        )

    )


# ==========================================================
# WORD DOCUMENTS
# ==========================================================

def compare_word(old_file, new_file):
    """
    Compare Word documents paragraph by paragraph.
    """

    old_doc = Document(old_file)

    new_doc = Document(new_file)

    old_text = [

        p.text.strip()

        for p in old_doc.paragraphs

        if p.text.strip()

    ]

    new_text = [

        p.text.strip()

        for p in new_doc.paragraphs

        if p.text.strip()

    ]

    return list(

        difflib.unified_diff(

            old_text,

            new_text,

            fromfile="Baseline Version",

            tofile="Current Version",

            lineterm=""

        )

    )


# ==========================================================
# EXCEL
# ==========================================================

def compare_excel(old_file, new_file):

    differences = []

    old_book = load_workbook(old_file)

    new_book = load_workbook(new_file)

    all_sheets = set(old_book.sheetnames) | set(new_book.sheetnames)

    for sheet_name in all_sheets:

        if sheet_name not in old_book.sheetnames:

            differences.append(
                f"[NEW SHEET] {sheet_name}"
            )

            continue

        if sheet_name not in new_book.sheetnames:

            differences.append(
                f"[DELETED SHEET] {sheet_name}"
            )

            continue

        old_sheet = old_book[sheet_name]

        new_sheet = new_book[sheet_name]

        rows = max(
            old_sheet.max_row,
            new_sheet.max_row
        )

        cols = max(
            old_sheet.max_column,
            new_sheet.max_column
        )

        for row in range(1, rows + 1):

            for col in range(1, cols + 1):

                old_value = old_sheet.cell(row, col).value

                new_value = new_sheet.cell(row, col).value

                if old_value != new_value:

                    cell = old_sheet.cell(row, col).coordinate

                    differences.append(

                        f"{sheet_name} | {cell} | "

                        f"{old_value} → {new_value}"

                    )

    return differences


# ==========================================================
# POWERPOINT
# ==========================================================

def extract_slide_text(slide):

    text = []

    for shape in slide.shapes:

        if hasattr(shape, "text"):

            if shape.text.strip():

                text.append(shape.text.strip())

    return text


def compare_powerpoint(old_file, new_file):

    differences = []

    old_ppt = Presentation(old_file)

    new_ppt = Presentation(new_file)

    total = max(

        len(old_ppt.slides),

        len(new_ppt.slides)

    )

    for index in range(total):

        if index >= len(old_ppt.slides):

            differences.append(
                f"[NEW SLIDE] {index + 1}"
            )

            continue

        if index >= len(new_ppt.slides):

            differences.append(
                f"[DELETED SLIDE] {index + 1}"
            )

            continue

        diff = list(

            difflib.unified_diff(

                extract_slide_text(
                    old_ppt.slides[index]
                ),

                extract_slide_text(
                    new_ppt.slides[index]
                ),

                fromfile=f"Slide {index+1}",

                tofile=f"Slide {index+1}",

                lineterm=""

            )

        )

        differences.extend(diff)

    return differences